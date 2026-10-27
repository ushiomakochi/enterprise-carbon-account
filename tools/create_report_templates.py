"""
报告模板生成器 — 一键创建所有 .docx 模板。

用法:
    python tools/create_report_templates.py

docxtpl 模板语法说明:
- {{ variable }}       — 简单变量替换
- {% for item in list %}...{% endfor %}  — 段落级循环
- {%tr for item in list %}               — 表格行循环（独占一行）
- {%tr endfor %}                         — 表格行循环结束（独占一行）
- {% if condition %}...{% endif %}       — 条件块

重要: {%tr for %} 和 {%tr endfor %} 必须各自独占一个独立表格行！
      它们所在的行会被 docxtpl 自动移除，中间的行被循环克隆。

模板列表:
    - verification.docx   核查报告
    - pdd.docx            项目设计文档（PDD）
"""

import os
import sys

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, PROJECT_ROOT)

OUTPUT_DIR = os.path.join(PROJECT_ROOT, "templates", "reports", "docx")


# ══════════════════════════════════════════════════════════════
# 辅助函数
# ══════════════════════════════════════════════════════════════

def add_dynamic_table(doc, headers, loop_var, loop_list, cell_texts, style="Table Grid"):
    """
    创建一个 docxtpl 动态表格（行数由数据驱动）。

    生成的表格结构:
        Row 0: 表头
        Row 1: {%tr for item in list %}  （docxtpl 会自动移除此行）
        Row 2: 模板行（每列填 {{ item.xxx }}）
        Row 3: {%tr endfor %}           （docxtpl 会自动移除此行）

    参数:
        doc: python-docx Document 对象
        headers: 表头列表，如 ["序号", "名称", "数值"]
        loop_var: 循环变量名，如 "src"
        loop_list: 循环的列表表达式，如 "emission_sources"
        cell_texts: 模板行每列的文本，如 ["{{ src.seq }}", "{{ src.name }}", ...]
        style: 表格样式

    返回:
        docx Table 对象
    """
    n_cols = len(headers)
    table = doc.add_table(rows=4, cols=n_cols, style=style)

    # Row 0: 表头
    for j, h in enumerate(headers):
        table.cell(0, j).text = h

    # Row 1: {%tr for %} — 独占一行，docxtpl 渲染时会移除
    table.cell(1, 0).text = f"{{%tr for {loop_var} in {loop_list} %}}"

    # Row 2: 模板行 — 会被循环克隆
    for j, text in enumerate(cell_texts):
        table.cell(2, j).text = text

    # Row 3: {%tr endfor %} — 独占一行，docxtpl 渲染时会移除
    table.cell(3, 0).text = "{%tr endfor %}"

    return table


def add_kv_table(doc, rows, style="Table Grid"):
    """创建一个静态键值对表格。

    参数:
        doc: Document
        rows: [(label, value), ...] 列表，第一行为表头
        style: 表格样式
    """
    table = doc.add_table(rows=len(rows), cols=2, style=style)
    for i, (label, value) in enumerate(rows):
        table.cell(i, 0).text = label
        table.cell(i, 1).text = value
    return table


# ══════════════════════════════════════════════════════════════
# 模板创建函数
# ══════════════════════════════════════════════════════════════

def create_verification_template():
    """创建核查报告模板"""
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT

    doc = Document()

    # ── 页面设置 ──
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.18)
    section.right_margin = Cm(3.18)

    style = doc.styles['Normal']
    style.font.name = 'SimSun'
    style.font.size = Pt(12)

    # ══ 封面 ══
    for _ in range(4):
        doc.add_paragraph()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("碳减排量核查报告")
    run.font.size = Pt(26)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x2D, 0x5A, 0x27)

    doc.add_paragraph()
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run("项目名称：{{ project.name }}\n").font.size = Pt(14)
    sub.add_run("项目编号：{{ project.code }}\n").font.size = Pt(14)
    sub.add_run("企业名称：{{ enterprise.name }}\n").font.size = Pt(14)

    doc.add_paragraph()
    date_p = doc.add_paragraph()
    date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_p.add_run("核查期间：{{ report.start_date }} 至 {{ report.end_date }}\n").font.size = Pt(12)
    date_p.add_run("报告生成时间：{{ report.generated_at }}").font.size = Pt(12)
    doc.add_page_break()

    # ══ 一、项目概况 ══
    doc.add_heading("一、项目概况", level=1)
    add_kv_table(doc, [
        ("项目名称", "{{ project.name }}"),
        ("项目编号", "{{ project.code }}"),
        ("项目类型", "{{ project.type }}"),
        ("所属公式", "{{ project.hook_label }}"),
        ("基准日期", "{{ project.baseline_date }}"),
        ("项目简介", "{{ project.summary }}"),
    ])
    doc.add_paragraph()

    # ══ 二、碳减排量核算 ══
    doc.add_heading("二、碳减排量核算", level=1)

    doc.add_heading("2.1 监测数据摘要", level=2)
    add_kv_table(doc, [
        ("指标", "数值"),
        ("监测记录总数", "{{ monitoring.total_records }} 条"),
        ("核查期总减排量", "{{ monitoring.total_reduction_fmt }} tCO₂e"),
        ("日均减排量", "{{ monitoring.daily_avg_fmt }} tCO₂e"),
        ("预估年减排量", "{{ monitoring.annual_estimate_fmt }} tCO₂e"),
    ])
    doc.add_paragraph()

    # 2.2 碳排放源清单（★ 动态表格 ★）
    doc.add_heading("2.2 碳排放源清单", level=2)
    doc.add_paragraph("以下清单列出了本核查期内项目的各项碳排放源及其对应的减排量。")

    add_dynamic_table(doc,
        headers=["序号", "排放源", "活动数据", "排放因子", "减排量(tCO₂e)"],
        loop_var="src", loop_list="emission_sources",
        cell_texts=[
            "{{ src.seq }}",
            "{{ src.name }}",
            "{{ src.activity_data }}",
            "{{ src.emission_factor }}",
            "{{ src.emission_amount }}",
        ])

    doc.add_paragraph("合计减排量：{{ monitoring.total_reduction_fmt }} tCO₂e")

    # 2.3 计算分解
    doc.add_heading("2.3 计算分解", level=2)
    add_dynamic_table(doc,
        headers=["参数", "数值"],
        loop_var="item", loop_list="monitoring.breakdown_items",
        cell_texts=["{{ item.key }}", "{{ item.value }}"])

    # ══ 三、计算方法 ══
    doc.add_heading("三、计算方法", level=1)

    doc.add_heading("3.1 计算公式", level=2)
    doc.add_paragraph("{{ formula.text }}")

    doc.add_heading("3.2 关键参数", level=2)
    add_dynamic_table(doc,
        headers=["参数名", "参数值"],
        loop_var="p", loop_list="formula.parameters.items()",
        cell_texts=["{{ p[0] }}", "{{ p[1] }}"])

    doc.add_paragraph("电网排放因子：{{ formula.ef_grid }} tCO₂e/MWh")

    # ══ 四、碳资产概况 ══
    doc.add_heading("四、碳资产概况", level=1)
    add_kv_table(doc, [
        ("指标", "数值 (tCO₂e)"),
        ("总碳资产", "{{ carbon_assets.total }}"),
        ("可用碳资产 (A)", "{{ carbon_assets.available_gross }}"),
        ("可用碳资产 (B)", "{{ carbon_assets.available_net }}"),
        ("已交易量", "{{ carbon_assets.traded }}"),
        ("累计核证减排量", "{{ carbon_assets.cumulative_reduction }}"),
    ])
    doc.add_paragraph()

    # ══ 五、项目阶段状态 ══
    doc.add_heading("五、项目阶段状态", level=1)
    add_dynamic_table(doc,
        headers=["阶段编号", "阶段名称", "审核状态"],
        loop_var="s", loop_list="stages",
        cell_texts=["阶段 {{ s.num }}", "{{ s.name }}", "{{ s.status }}"])
    doc.add_paragraph()

    # ══ 六、结论与声明 ══
    doc.add_heading("六、结论", level=1)
    p = doc.add_paragraph()
    p.add_run("{% if monitoring.total_reduction > 0 %}")
    p.add_run("基于上述监测数据和计算方法，本核查期内项目碳减排量核算结果真实、准确。项目运行正常，减排效果显著。")
    p.add_run("{% else %}")
    p.add_run("本核查期内暂无监测数据，建议启动数据采集后重新生成报告。")
    p.add_run("{% endif %}")

    doc.add_heading("七、声明", level=1)
    doc.add_paragraph("{% for d in declarations %}{{ d }}{% endfor %}")

    doc.add_paragraph()
    fp = doc.add_paragraph()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.add_run("{{ report.generator }} | {{ report.generated_at }}").font.size = Pt(9)

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    out = os.path.join(OUTPUT_DIR, "verification.docx")
    doc.save(out)
    print(f"  [OK] verification: {out}")


def create_pdd_template():
    """创建 PDD 项目设计文档模板"""
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)

    style = doc.styles['Normal']
    style.font.name = 'SimSun'
    style.font.size = Pt(12)

    # ══ 封面 ══
    for _ in range(4):
        doc.add_paragraph()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("项目设计文档（PDD）")
    run.font.size = Pt(24)
    run.font.bold = True

    doc.add_paragraph()
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run("Project Design Document\n\n自愿型碳减排项目（VCER）\n").font.size = Pt(12)

    doc.add_paragraph()
    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info.add_run("项目名称：{{ project.name }}\n").font.size = Pt(14)
    info.add_run("项目编号：{{ project.code }}\n").font.size = Pt(14)
    info.add_run("企业名称：{{ enterprise.name }}\n").font.size = Pt(14)
    info.add_run("报告生成时间：{{ report.generated_at }}").font.size = Pt(12)
    doc.add_page_break()

    # ══ A. 项目概述 ══
    doc.add_heading("A. 项目概述", level=1)

    doc.add_heading("A.1 项目基本信息", level=2)
    add_kv_table(doc, [
        ("项目名称", "{{ project.name }}"),
        ("项目类型", "{{ pdd.overview.project_type }}"),
        ("项目编号", "{{ pdd.overview.project_code }}"),
        ("计入期类型", "{{ pdd.crediting_period.type }}"),
        ("首个计入期", "{{ pdd.crediting_period.first_period }}"),
        ("计入期起始日", "{{ pdd.crediting_period.start_date }}"),
    ])
    doc.add_paragraph()

    doc.add_heading("A.2 基准线情景", level=2)
    doc.add_paragraph("{{ pdd.overview.baseline_scenario }}")
    doc.add_heading("A.3 项目情景", level=2)
    doc.add_paragraph("{{ pdd.overview.project_scenario }}")

    # ══ B. 方法学 ══
    doc.add_heading("B. 方法学", level=1)
    doc.add_heading("B.1 适用方法学", level=2)
    doc.add_paragraph("{{ pdd.methodology.reference }}")
    doc.add_heading("B.2 适用性分析", level=2)
    doc.add_paragraph("{{ pdd.methodology.applicability }}")
    doc.add_heading("B.3 计算公式", level=2)
    doc.add_paragraph("{{ pdd.methodology.formula_text }}")

    doc.add_heading("B.4 关键参数", level=2)
    add_dynamic_table(doc,
        headers=["参数名", "参数值"],
        loop_var="p", loop_list="pdd.methodology.parameters",
        cell_texts=["{{ p.key }}", "{{ p.value }}"])

    # ══ C. 额外性论证 ══
    doc.add_heading("C. 额外性论证", level=1)
    doc.add_heading("C.1 投资分析", level=2)
    doc.add_paragraph("{{ pdd.additionality.investment_analysis }}")
    doc.add_heading("C.2 屏障分析", level=2)
    doc.add_paragraph("{{ pdd.additionality.barrier_analysis }}")
    doc.add_heading("C.3 普遍性分析", level=2)
    doc.add_paragraph("{{ pdd.additionality.common_practice }}")

    # ══ D. 减排量估算 ══
    doc.add_heading("D. 减排量估算", level=1)

    doc.add_heading("D.1 实际已产生的温室气体减排量", level=2)
    doc.add_paragraph("统计期间：{{ report.start_date }} 至 {{ report.end_date }}")
    add_kv_table(doc, [
        ("指标", "数值"),
        ("期内实际减排量", "{{ monitoring.total_reduction_fmt }} tCO₂e"),
        ("期内基准线排放", "{{ monitoring.baseline_emissions_fmt }} tCO₂e"),
        ("期内项目排放", "{{ monitoring.project_emissions_fmt }} tCO₂e"),
    ])
    doc.add_paragraph()

    doc.add_heading("D.2 计入期逐年预估", level=2)
    doc.add_paragraph("以下为计入期内逐年预估的基准线排放（表 B.2，由全项目通用预设 project_preset.ten_year_baseline 提供）。")
    doc.add_paragraph("{{ project_preset.ten_year_baseline }}")
    doc.add_paragraph("以下为计入期内逐年预估的项目排放（表 B.3，由全项目通用预设 project_preset.ten_year_project_emission 提供）。")
    doc.add_paragraph("{{ project_preset.ten_year_project_emission }}")
    doc.add_paragraph("以下为计入期内逐年预估的基准线排放、项目排放、泄漏与碳减排量（表 B.4，由全项目通用预设 project_preset.ten_year_projection 提供）。")
    doc.add_paragraph("{{ project_preset.ten_year_projection }}")
    doc.add_paragraph()

    doc.add_heading("D.3 碳排放源清单", level=2)
    add_dynamic_table(doc,
        headers=["序号", "排放源", "活动数据", "排放因子", "减排量(tCO₂e)"],
        loop_var="src", loop_list="emission_sources",
        cell_texts=[
            "{{ src.seq }}", "{{ src.name }}",
            "{{ src.activity_data }}", "{{ src.emission_factor }}",
            "{{ src.emission_amount }}",
        ])
    doc.add_paragraph()

    # ══ E. 环境影响评价 ══
    doc.add_heading("E. 环境影响评价", level=1)
    doc.add_paragraph("{{ pdd.environmental_impact }}")

    # ══ F. 利益相关方意见 ══
    doc.add_heading("F. 利益相关方意见", level=1)
    doc.add_paragraph("{{ pdd.stakeholder_comments }}")

    # ══ G. 监测计划 ══
    doc.add_heading("G. 监测计划", level=1)
    doc.add_heading("G.1 监测参数", level=2)
    add_dynamic_table(doc,
        headers=["参数", "单位", "监测频率", "监测方法"],
        loop_var="mp", loop_list="pdd.monitoring_plan.parameters",
        cell_texts=[
            "{{ mp.parameter }}", "{{ mp.unit }}",
            "{{ mp.frequency }}", "{{ mp.method }}",
        ])
    doc.add_paragraph()

    doc.add_heading("G.2 质量保证/质量控制", level=2)
    doc.add_paragraph("{{ pdd.monitoring_plan.qa_qc }}")

    # ══ H. 计入期 ══
    doc.add_heading("H. 计入期", level=1)
    add_kv_table(doc, [
        ("计入期类型", "{{ pdd.crediting_period.type }}"),
        ("首个计入期", "{{ pdd.crediting_period.first_period }}"),
        ("最大更新次数", "{{ pdd.crediting_period.max_periods }}"),
        ("起始日期", "{{ pdd.crediting_period.start_date }}"),
    ])
    doc.add_paragraph()

    doc.add_paragraph()
    fp = doc.add_paragraph()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.add_run("{{ report.generator }} | {{ report.generated_at }}").font.size = Pt(9)

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    out = os.path.join(OUTPUT_DIR, "pdd.docx")
    doc.save(out)
    print(f"  [OK] pdd: {out}")


# ══════════════════════════════════════════════════════════════
# 主入口
# ══════════════════════════════════════════════════════════════

def main():
    print("=" * 60)
    print("  Carbon Report Word Template Generator")
    print("  docxtpl + Jinja2 .docx templates")
    print("=" * 60)
    print()

    os.makedirs(OUTPUT_DIR, exist_ok=True)

    for fn in [create_verification_template, create_pdd_template]:
        try:
            fn()
        except Exception as e:
            import traceback
            print(f"  [ERROR] {fn.__name__}: {e}")
            traceback.print_exc()

    print()
    print("=" * 60)
    print("  Templates created!")
    print(f"  Output: {OUTPUT_DIR}")
    print()
    print("  Next steps:")
    print("  1. Open templates in Word to adjust layout/styles")
    print("  2. Add company logo, headers, footers")
    print("  3. Don't modify {{ }} and {% %} tags (Jinja2 syntax)")
    print("  4. Add new .docx files to this dir for auto-discovery")
    print("=" * 60)


if __name__ == "__main__":
    main()
