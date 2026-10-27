"""
一键生成企业报告模板 — 基于参考报告，将动态数据替换为 Jinja2 占位符。

用法:
    python tools/_create_enterprise_template.py ^
        --reference "参考报告.docx" --enterprise "<企业名>" --report verification

说明:
    这是适配新企业时的**一次性工具**：占位符替换映射表是按参考报告的
    实际内容写的，每家企业都要按自己的报告改映射表（见 main() 中 TODO）。
    模板生成后此工具不会再被调用 —— 系统运行时只认生成的 .docx。

依赖: python-docx (已在项目中使用)
"""
import argparse
import os
import re
import sys
import shutil

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, PROJECT_ROOT)


def replace_text_in_docx(doc, replacements):
    """
    在 docx 文档的段落和表格中进行全文文本替换。

    策略（简化但可靠）：
    1. 获取段落/单元格完整文本
    2. 对所有匹配项做字符串替换
    3. 如果有改动：将结果写入第一个 run，清空其余 run
    （这会丢失部分行内格式，但确保替换可靠）
    """
    import re

    # 按 key 长度降序排列，防止短 key 先匹配导致长 key 中的短子串被替换
    sorted_items = sorted(replacements.items(), key=lambda x: len(x[0]), reverse=True)

    def apply_to_paragraph(paragraph):
        runs = paragraph.runs
        if not runs:
            return

        original = paragraph.text
        if not original:
            return

        # 检查是否有任何匹配
        new_text = original
        for old, new in sorted_items:
            new_text = new_text.replace(old, new)

        if new_text == original:
            return  # 没有改动

        # 有改动：保留第一个 run 的格式，写入全部文本
        runs[0].text = new_text
        for run in runs[1:]:
            run.text = ""

    # 处理所有段落
    for paragraph in doc.paragraphs:
        apply_to_paragraph(paragraph)

    # 处理所有表格
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    apply_to_paragraph(paragraph)

    # 处理页眉页脚
    for section in doc.sections:
        for paragraph in section.header.paragraphs:
            apply_to_paragraph(paragraph)
        for paragraph in section.footer.paragraphs:
            apply_to_paragraph(paragraph)

    return doc


def main():
    parser = argparse.ArgumentParser(description="基于参考报告生成企业 docx 模板")
    parser.add_argument("--reference", required=True, help="参考报告 .docx 路径")
    parser.add_argument("--enterprise", required=True, help="企业名（模板放到 docx/<企业名>/ 下）")
    parser.add_argument("--report", default="verification",
                        help="报告类型 key：verification / pdd / approval（默认 verification）")
    args = parser.parse_args()

    reference_docx = args.reference
    output_dir = os.path.join(PROJECT_ROOT, "templates", "reports", "docx", args.enterprise)
    output_file = os.path.join(output_dir, f"{args.report}.docx")

    print("=" * 60)
    print(f"  创建企业报告模板: {args.enterprise}")
    print(f"  报告类型: {args.report}")
    print(f"  参考报告: {os.path.basename(reference_docx)}")
    print("=" * 60)

    if not os.path.exists(reference_docx):
        print(f"[ERROR] 参考报告不存在: {reference_docx}")
        sys.exit(1)

    os.makedirs(output_dir, exist_ok=True)

    # 先复制参考报告
    print(f"\n[1/2] 复制参考报告...")
    shutil.copy2(reference_docx, output_file)
    print(f"  → {output_file}")

    # 打开并替换
    print(f"[2/2] 替换动态数据为 Jinja2 占位符...")
    from docx import Document
    doc = Document(output_file)

    # ── 替换映射表 ──
    # 格式: "原始值" → "{{ placeholder }}"
    #
    # 注意事项:
    # 1. 数值在全文中有多处出现，必须全部替换（包括表格和段落）
    # 2. 替换顺序很重要：较长的文本先替换，避免部分匹配
    # 3. 占位符命名遵循 context 结构（参考 context_builder.py）
    #
    # ⚠️ TODO: 下面的映射表是**示例**（来自首个接入企业的核查报告），
    #    给新企业生成模板时，必须按该企业的参考报告实际内容逐项改写。

    replacements = {
        # ── 日期 ──
        "2025年1月1日-2025年12月31日": "{{ report.start_date }}-{{ report.end_date }}",
        "2025年1月1日00:00-2025年12月31日23:59": "{{ report.start_date }} 00:00-{{ report.end_date }} 23:59",
        "2026年6月14日": "{{ report.generated_at }}",
        # 监测频次的日期（cell 内每个换行是独立 paragraph，需分开替换）
        "2025年1月1日00:00一次（开始）": "{{ report.start_date }} 00:00一次（开始）",
        "2025年12月31日23:59一次（结束）": "{{ report.end_date }} 23:59一次（结束）",

        # ── 封面表 & 减排量汇总 → 关键数值 ──
        # 注意: 302.39 和 279.96 在封面上和汇总表中都有
        "302.39": "{{ monitoring.pre_estimated_reduction_fmt }}",
        "279.96": "{{ monitoring.total_reduction_fmt }}",
        "835.84": "{{ monitoring.baseline_emissions_fmt }}",
        "555.88": "{{ monitoring.project_emissions_fmt }}",

        # ── 监测参数表 → 监测数值 ──
        "962222.2": "{{ monitoring.consumed_power }}",
        "8826147": "{{ monitoring.compressed_air_volume }}",
        "525599": "{{ monitoring.runtime }}",

        # ── P84 段落中的参数 ──
        "额定排气压力0.7": "额定排气压力{{ project.rated_pressure }}",
        "电网排放因子0.5777": "电网排放因子{{ formula.ef_grid }}",
        "核算时长525599": "核算时长{{ monitoring.runtime }}",

    }

    replace_text_in_docx(doc, replacements)
    doc.save(output_file)
    print(f"  → 已保存: {output_file}")

    # ── 验证 ──
    print(f"\n[验证] 检查替换结果...")
    doc2 = Document(output_file)
    found_placeholders = set()
    # 匹配 {{ var.name }}（允许空格）
    placeholder_re = re.compile(r'\{\{\s*([a-zA-Z_][\w.]*)\s*\}\}')
    for p in doc2.paragraphs:
        found_placeholders.update(placeholder_re.findall(p.text))
    for table in doc2.tables:
        for row in table.rows:
            for cell in row.cells:
                found_placeholders.update(placeholder_re.findall(cell.text))

    print(f"  模板中的占位符 ({len(found_placeholders)} 个):")
    for ph in sorted(found_placeholders):
        print(f"    - {{{{ {ph} }}}}")

    print()
    print("=" * 60)
    print("  模板创建完成!")
    print(f"  文件: {output_file}")
    print()
    print("  说明: 固定文本（企业简介、项目描述、方法学）已保留在模板中。")
    print("  生成报告时，系统会自动填入动态数据。")
    print("=" * 60)


if __name__ == "__main__":
    main()
